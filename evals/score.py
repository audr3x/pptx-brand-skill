#!/usr/bin/env python3
"""Score eval runs and write evals/scorecard.md.

Usage:
  python3 evals/score.py                      # all models under evals/runs/
  python3 evals/score.py --model claude-opus-5.5
  python3 evals/score.py --packet claude-opus-5.5   # write the visual judge packet

Each task gets seven automated category scores from 0 to 1 (see evals/RUBRIC.md), plus a
judge score when evals/runs/<model>/<task>/judge.json exists. A category that does not apply
(for example the build on a task where asking questions is the right answer) is left out.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any

import yaml

EVALS = Path(__file__).resolve().parent
ROOT = EVALS.parent
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(EVALS))

from checks import (  # noqa: E402
    chunk_problems,
    load_requests,
    pptx_problems,
    pptx_text_by_slide,
    slides_plan_problems,
    slides_text_by_slide,
)
from outline_rules import LAYOUT_ALIASES, canonical_layout, prepare  # noqa: E402

CATEGORIES = ["structure", "brand", "layouts", "faithful", "instructions", "render", "process"]
NUMBER = re.compile(r"(?<![\w.])\$?\d[\d,]*(?:\.\d+)?%?")


def norm(text: Any) -> str:
    text = str(text).replace("\u2019", "'").replace("\u201c", '"').replace("\u201d", '"')
    return re.sub(r"\s+", " ", text).strip().lower()


def numbers(text: str) -> set[float]:
    out = set()
    for tok in NUMBER.findall(text):
        plain = tok.strip("$%").replace(",", "")
        try:
            value = float(plain)
        except ValueError:
            continue
        if "." in plain or "%" in tok or "$" in tok or abs(value) >= 10:
            out.add(value)
    return out


def all_numbers(text: str) -> set[float]:
    out = set()
    for tok in NUMBER.findall(text):
        try:
            out.add(float(tok.strip("$%").replace(",", "")))
        except ValueError:
            pass
    return out


def get_path(spec: Any, path: str) -> Any:
    cur = spec
    for part in path.split("."):
        if isinstance(cur, list):
            try:
                cur = cur[int(part)]
            except (ValueError, IndexError):
                return None
        elif isinstance(cur, dict):
            cur = cur.get(part)
        else:
            return None
    return cur


def chart_series(outline: dict) -> list[tuple[str, list[float]]]:
    out = []
    for s in outline.get("slides", []):
        chart = s.get("chart") if isinstance(s, dict) else None
        if isinstance(chart, dict):
            for ser in chart.get("series") or []:
                if isinstance(ser, dict):
                    out.append((str(ser.get("name", "")), [float(v) for v in ser.get("values") or [] if isinstance(v, (int, float))]))
    return out


def alias(layout: str) -> str:
    """Canonical template layout for an outline alias like 'two-column'."""
    return LAYOUT_ALIASES.get(layout.lower(), layout)


def mean(values: list[float]) -> float | None:
    return sum(values) / len(values) if values else None


def score_task(model: str, task: str) -> dict[str, Any]:
    spec = yaml.safe_load((EVALS / "specs" / f"{task}.yaml").read_text())
    prompt = (EVALS / "tasks" / f"{task}.md").read_text()
    run_dir = EVALS / "runs" / model / task
    target = spec.get("target", "pptx")
    row: dict[str, Any] = {"task": task, "difficulty": spec.get("difficulty", ""), "notes": []}
    scores: dict[str, float | None] = {c: None for c in CATEGORIES}
    response = (run_dir / "response.md").read_text() if (run_dir / "response.md").is_file() else ""
    outline_path = run_dir / "outline.yaml"
    build_required = spec.get("build_required", True)

    outline = None
    errors: list[str] = []
    warnings: list[str] = []
    if outline_path.is_file():
        try:
            outline = yaml.safe_load(outline_path.read_text())
            outline, errors, warnings = prepare(outline, target, "outline")
        except yaml.YAMLError:
            errors = ["outline.yaml is not valid YAML"]
            outline = None
    elif build_required:
        row["notes"].append("no outline.yaml")

    # Built deck text and artifacts
    slide_texts: list[str] = []
    built = False
    brand_hits: list[str] = []
    if target == "pptx" and (run_dir / "deck.pptx").is_file():
        built = True
        slide_texts = pptx_text_by_slide(run_dir / "deck.pptx")
        problems = pptx_problems(run_dir / "deck.pptx")
        brand_hits = [f"{k}: {v[0]}" for k, v in problems.items() if v]
        brand_total = len(problems)
    elif target == "slides" and (run_dir / "slides" / "fills.json").is_file() and (run_dir / "slides" / "creates.json").is_file():
        built = True
        reqs = load_requests(run_dir / "slides" / "creates.json") + load_requests(run_dir / "slides" / "fills.json")
        slide_texts = slides_text_by_slide(reqs)
        problems = slides_plan_problems(reqs)
        problems["chunks"] = chunk_problems(run_dir / "slides" / "fills.json")
        brand_hits = [f"{k}: {v[0]}" for k, v in problems.items() if v]
        brand_total = len(problems)
        # Slides draws chart values as bar heights, not text.
        if isinstance(outline, dict):
            slide_texts.append(" ".join(f"{v:g}" for _, vals in chart_series(outline) for v in vals))
    deck_text = norm(" ".join(slide_texts))
    if not deck_text and isinstance(outline, dict):
        deck_text = norm(" ".join(str(v) for v in _leaves(outline)))

    # structure
    if outline is not None or build_required:
        scores["structure"] = mean([float(outline is not None and not errors), float(built)]) if build_required else float(not errors)
    if errors:
        row["notes"].append(f"check errors: {errors[0][:80]}")

    # brand
    if built:
        scores["brand"] = 1 - len(brand_hits) / brand_total
        row["notes"].extend(brand_hits[:2])

    # layouts and slide count
    if isinstance(outline, dict) and isinstance(outline.get("slides"), list):
        got = [canonical_layout(s) if isinstance(s, dict) else "" for s in outline["slides"]]
        parts = []
        if "slides" in spec:
            lo, hi = spec["slides"]
            parts.append(float(lo <= len(got) <= hi))
            if not parts[-1]:
                row["notes"].append(f"{len(got)} slides, expected {lo} to {hi}")
        if "layouts" in spec:
            want = [alias(l) for l in spec["layouts"]]
            parts.append(sum(a == b for a, b in zip(want, got)) / max(len(want), len(got)))
        if "layouts_include" in spec:
            want = [alias(l) for l in spec["layouts_include"]]
            parts.append(sum(w in got for w in want) / len(want))
        scores["layouts"] = mean(parts)

    # faithfulness
    if deck_text:
        parts = []
        facts = spec.get("facts") or []
        if facts:
            missing = [f for f in facts if norm(f) not in deck_text]
            parts.append(1 - len(missing) / len(facts))
            if missing:
                row["notes"].append(f"missing facts: {', '.join(missing[:3])}")
        allowed = all_numbers(prompt) | {float(n) for n in spec.get("allowed_numbers", [])}
        invented = sorted(numbers(deck_text) - allowed)
        parts.append(max(0.0, 1 - 0.25 * len(invented)))
        if invented:
            row["notes"].append(f"numbers not in the prompt: {', '.join(f'{n:g}' for n in invented[:5])}")
        scores["faithful"] = mean(parts)

    # instructions
    results = []
    for check in spec.get("checks") or []:
        (kind, arg), = check.items()
        ok, why = run_check(kind, arg, outline, run_dir, response, warnings)
        results.append(float(ok))
        if not ok:
            row["notes"].append(why)
    if deck_text:
        dashes = "\u2014" in " ".join(slide_texts) or "\u2013" in " ".join(slide_texts)
        results.append(float(not dashes))
        if dashes:
            row["notes"].append("em or en dash in deck text")
    scores["instructions"] = mean(results)

    # render
    attempts = [json.loads(l) for l in (run_dir / "attempts.jsonl").read_text().splitlines()] if (run_dir / "attempts.jsonl").is_file() else []
    pngs = list((run_dir / "render").glob("slide-*.png")) if (run_dir / "render").is_dir() else []
    if build_required or pngs:
        expected = len(outline.get("slides", [])) if isinstance(outline, dict) else 0
        scores["render"] = float(bool(pngs) and len(pngs) == expected)
        if not scores["render"]:
            row["notes"].append(f"{len(pngs)} PNGs for {expected} slides")

    # process
    failed = sum(1 for a in attempts if a.get("rc"))
    if attempts or build_required:
        tracebacks = any(s.get("traceback") for a in attempts for s in a.get("steps", []))
        scores["process"] = 0.0 if not attempts else max(0.0, 1 - 0.25 * failed) * (0.5 if tracebacks else 1)
    row["attempts"] = len(attempts)
    row["failed_attempts"] = failed

    judge = run_dir / "judge.json"
    if judge.is_file():
        j = json.loads(judge.read_text())
        vals = [j[k] for k in ("visual", "content", "communication") if isinstance(j.get(k), (int, float))]
        row["judge"] = (mean(vals) - 1) / 4 if vals else None
    else:
        row["judge"] = None
    row["scores"] = scores
    auto = [v for v in scores.values() if v is not None]
    row["auto"] = mean(auto)
    return row


def _leaves(value: Any):
    if isinstance(value, dict):
        for k, v in value.items():
            if k not in {"layout", "intent", "image"}:
                yield from _leaves(v)
    elif isinstance(value, list):
        for v in value:
            yield from _leaves(v)
    elif value is not None:
        yield value


def run_check(kind: str, arg: Any, outline: Any, run_dir: Path, response: str, warnings: list[str]) -> tuple[bool, str]:
    slides = outline.get("slides", []) if isinstance(outline, dict) else []
    resp = norm(response)
    if kind == "field_equals":
        idx = arg["slide"] - 1
        val = get_path(slides[idx], arg["path"]) if idx < len(slides) else None
        if arg["path"].startswith("bullets.") and isinstance(val, dict):
            val = val.get("text")
        return norm(val or "") == norm(arg["value"]), f"slide {arg['slide']} {arg['path']} is {str(val)[:40]!r}"
    if kind == "chart_values":
        want = [float(v) for v in arg["values"]]
        series = chart_series(outline) if isinstance(outline, dict) else []
        ok = any((arg["series"] == "any" or norm(n) == norm(arg["series"])) and vals == want for n, vals in series)
        return ok, f"chart series {arg['series']} values do not match {arg['values']}"
    if kind == "chart":
        charts = [s["chart"] for s in slides if isinstance(s, dict) and isinstance(s.get("chart"), dict)]
        for c in charts:
            ser = c.get("series") or []
            ok = (str(c.get("type", "column")) == arg.get("type", c.get("type", "column"))
                  and ("series" not in arg or len(ser) == arg["series"])
                  and ("categories" not in arg or len(c.get("categories") or []) == arg["categories"])
                  and (not arg.get("negative") or any(v < 0 for s in ser for v in s.get("values", []) if isinstance(v, (int, float)))))
            if ok:
                return True, ""
        return False, f"no chart matching {arg}"
    if kind == "table_rows":
        rows = [len(s.get("rows") or []) for s in slides if isinstance(s, dict) and s.get("headers")]
        return sum(rows) == arg, f"table rows {rows}, expected {arg}"
    if kind == "max_table_rows":
        rows = [len(s.get("rows") or []) for s in slides if isinstance(s, dict) and s.get("headers")]
        return bool(rows) and max(rows) <= arg, f"table rows per slide {rows}, max {arg}"
    if kind == "max_bullets":
        from outline_rules import count_bullets

        counts = [count_bullets(s.get("bullets") or []) for s in slides if isinstance(s, dict) and isinstance(s.get("bullets"), list)]
        return all(c <= arg for c in counts), f"bullets per slide {counts}, max {arg}"
    if kind == "no_title_warnings":
        hits = [w for w in warnings if "title is" in w or "eyebrow is" in w]
        return not hits, hits[0] if hits else ""
    if kind == "response_mentions":
        missing = [g for g in arg if not any(t in resp for t in g)]
        return not missing, f"response does not mention {missing}"
    if kind == "unchanged_from":
        base, _, _ = prepare(yaml.safe_load((ROOT / arg["path"]).read_text()), "pptx", "fixture")
        exc = {int(k): set(v) for k, v in (arg.get("except") or {}).items()}
        diffs = []
        for i, b in enumerate(base["slides"], 1):
            g = slides[i - 1] if i - 1 < len(slides) else {}
            for key in set(b) | set(g):
                if key in exc.get(i, set()):
                    continue
                if b.get(key) != g.get(key):
                    diffs.append(f"s{i}.{key}")
        if isinstance(outline, dict):
            for key in set(base) - {"slides"}:
                if base.get(key) != outline.get(key):
                    diffs.append(key)
        return not diffs, f"changed outside the request: {', '.join(diffs[:4])}"
    if kind == "fixture_untouched":
        import subprocess

        dirty = subprocess.run(["git", "status", "--porcelain", arg], cwd=ROOT, capture_output=True, text=True).stdout.strip()
        return not dirty, f"{arg} was edited"
    if kind == "todo_or_question":
        text = (outline_path_text(run_dir) + "\n" + response)
        return "TODO" in text or "?" in response, "no TODO in the outline and no question in the response"
    if kind == "no_images":
        return not any(isinstance(s, dict) and s.get("image") for s in slides), "image field on a no-image deck"
    if kind == "pictures":
        deck = run_dir / "deck.pptx"
        if not deck.is_file():
            return False, "no deck.pptx"
        from pptx import Presentation

        n = sum(1 for s in Presentation(str(deck)).slides for sh in s.shapes if sh._element.tag.endswith("}pic"))
        return n >= arg, f"{n} pictures, expected {arg}"
    return False, f"unknown check {kind}"


def outline_path_text(run_dir: Path) -> str:
    p = run_dir / "outline.yaml"
    return p.read_text() if p.is_file() else ""


def pct(v: float | None) -> str:
    return "n/a" if v is None else f"{round(v * 100)}"


def scorecard(models: list[str]) -> str:
    tasks = sorted(p.stem for p in (EVALS / "specs").glob("*.yaml"))
    lines = ["# Eval scorecard", "", "Generated by `python3 evals/score.py`. Scores are 0 to 100. See `evals/RUBRIC.md`.", ""]
    summary = ["## Summary", "", "| Model | Tasks run | Automated | Judge | Failed attempts | " + " | ".join(CATEGORIES) + " |",
               "| --- | --- | --- | --- | --- | " + " | ".join("---" for _ in CATEGORIES) + " |"]
    detail: list[str] = []
    for model in models:
        rows = [score_task(model, t) for t in tasks if (EVALS / "runs" / model / t).is_dir()]
        if not rows:
            continue
        cat = {c: mean([r["scores"][c] for r in rows if r["scores"][c] is not None]) for c in CATEGORIES}
        judge = mean([r["judge"] for r in rows if r["judge"] is not None])
        summary.append(f"| {model} | {len(rows)}/{len(tasks)} | {pct(mean([r['auto'] for r in rows if r['auto'] is not None]))} | {pct(judge)} | "
                       f"{sum(r['failed_attempts'] for r in rows)} | " + " | ".join(pct(cat[c]) for c in CATEGORIES) + " |")
        detail += [f"## {model}", "", "| Task | Level | Auto | Judge | Attempts | " + " | ".join(CATEGORIES) + " | Notes |",
                   "| --- | --- | --- | --- | --- | " + " | ".join("---" for _ in CATEGORIES) + " | --- |"]
        for r in rows:
            notes = "; ".join(r["notes"][:3]).replace("|", "/")[:200]
            detail.append(f"| {r['task']} | {r['difficulty']} | {pct(r['auto'])} | {pct(r['judge'])} | {r['attempts']} ({r['failed_attempts']} failed) | "
                          + " | ".join(pct(r["scores"][c]) for c in CATEGORIES) + f" | {notes} |")
        detail.append("")
    return "\n".join(lines + summary + [""] + detail)


def packet(model: str) -> Path:
    out = EVALS / "runs" / model / "judge-packet.md"
    lines = [f"# Judge packet: {model}", "", "Follow evals/JUDGE.md. For each task, read the prompt, the outline, the response, and every PNG, "
             "then write judge.json in the task folder.", ""]
    for spec in sorted((EVALS / "specs").glob("*.yaml")):
        task = spec.stem
        run_dir = EVALS / "runs" / model / task
        if not run_dir.is_dir():
            continue
        target = yaml.safe_load(spec.read_text()).get("target")
        pngs = sorted((run_dir / "render").glob("slide-*.png"), key=lambda p: int(p.stem.split("-")[1])) if (run_dir / "render").is_dir() else []
        lines += [f"## {task} ({target})", "", f"- Prompt: `evals/tasks/{task}.md`", f"- Outline: `{(run_dir / 'outline.yaml').relative_to(ROOT)}`",
                  f"- Response to the user: `{(run_dir / 'response.md').relative_to(ROOT)}`",
                  "- Slides: " + (", ".join(f"`{p.relative_to(ROOT)}`" for p in pngs) or "none rendered"),
                  f"- Write: `{(run_dir / 'judge.json').relative_to(ROOT)}`", ""]
        if target == "slides":
            lines += ["The PNGs are the PowerPoint built from the same outline. Judge layout and content from them; "
                      "Slides-only differences (column charts, no photos or notes) are expected.", ""]
    out.write_text("\n".join(lines))
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--model", action="append", help="score only these model folders")
    ap.add_argument("--packet", help="write the judge packet for this model folder and exit")
    ap.add_argument("--out", type=Path, default=EVALS / "scorecard.md")
    args = ap.parse_args()
    if args.packet:
        print(packet(args.packet).relative_to(ROOT))
        return 0
    runs = EVALS / "runs"
    models = args.model or (sorted(p.name for p in runs.iterdir() if p.is_dir() and not p.name.startswith("_")) if runs.is_dir() else [])
    text = scorecard(models)
    args.out.write_text(text + "\n")
    print(text)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
