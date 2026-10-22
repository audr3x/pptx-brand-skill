"""Deterministic checks on built decks. Used by evals/score.py and evals/adversarial/run.py.

Each check returns a list of problems (empty means pass).
"""
from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
TOKENS = json.loads((ROOT / "assets" / "brand-tokens.json").read_text())
BRAND_HEX = {v.upper() for v in {**TOKENS["core"], **TOKENS["secondary"]}.values()}
PPTX_FONTS = {"Archivo", "Archivo Light", "Arial"}
NO_FOOTER_LAYOUTS = {"Title Slide", "Bumper Slide"}
MAX_CHUNK = 50


# ---------- Google Slides request plans ----------

def _hex(rgb: dict[str, float]) -> str:
    return "".join(f"{round(rgb.get(c, 0) * 255):02X}" for c in ("red", "green", "blue"))


def _walk(value: Any):
    if isinstance(value, dict):
        yield value
        for v in value.values():
            yield from _walk(v)
    elif isinstance(value, list):
        for v in value:
            yield from _walk(v)


def slides_plan_problems(requests: list[dict[str, Any]]) -> dict[str, list[str]]:
    """Problems by check name for a Slides batchUpdate request list (creates + fills)."""
    out: dict[str, list[str]] = {"references": [], "fonts": [], "colors": [], "bullets": [], "images": []}
    known: set[str] = set()
    bulleted: dict[str, int] = {}
    for n, req in enumerate(requests):
        kind, body = next(iter(req.items()))
        if kind == "createSlide":
            known.add(body["objectId"])
            known.update(m["objectId"] for m in body.get("placeholderIdMappings", []))
            continue
        if kind in {"createShape", "createTable", "createImage", "createLine"}:
            page = body.get("elementProperties", {}).get("pageObjectId")
            if page and page not in known:
                out["references"].append(f"#{n} {kind} on unknown page {page}")
            known.add(body["objectId"])
            if kind == "createImage":
                url = body.get("url", "")
                if re.search(r"YOUR_|<|\.\.\.", url) or "raw.githubusercontent.com/audr3x/" in url:
                    out["images"].append(f"#{n} createImage url is a placeholder or private: {url[:80]}")
            continue
        targets = [body.get("objectId")] if body.get("objectId") else []
        targets += body.get("pageElementObjectIds", [])
        for t in targets:
            if t not in known:
                out["references"].append(f"#{n} {kind} targets {t} before it exists")
        if kind == "createParagraphBullets":
            bulleted[body["objectId"]] = n
        if kind == "updateParagraphStyle" and body.get("objectId") in bulleted:
            style = body.get("style", {})
            if style.get("spacingMode") == "NEVER_COLLAPSE" and "spaceBelow" in style:
                bulleted.pop(body["objectId"])
    for oid, n in bulleted.items():
        out["bullets"].append(f"#{n} {oid}: bullets without spaceAbove/spaceBelow and spacingMode NEVER_COLLAPSE after them")
    for node in _walk(requests):
        fam = node.get("fontFamily")
        if fam is not None and fam != "Archivo":
            out["fonts"].append(f"fontFamily {fam!r}")
        wff = node.get("weightedFontFamily")
        if isinstance(wff, dict) and (wff.get("fontFamily") != "Archivo" or wff.get("weight") not in {300, 400, 700}):
            out["fonts"].append(f"weightedFontFamily {wff}")
        rgb = node.get("rgbColor")
        if isinstance(rgb, dict) and _hex(rgb) not in BRAND_HEX:
            out["colors"].append(f"off-brand color #{_hex(rgb)}")
    return {k: sorted(set(v))[:10] for k, v in out.items()}


def slides_text_by_slide(requests: list[dict[str, Any]]) -> list[str]:
    """Text inserted on each wd_sN slide, without chart axis tick labels."""
    texts: dict[int, list[str]] = {}
    for req in requests:
        ins = req.get("insertText")
        if not ins or "_tick" in ins["objectId"]:
            continue
        m = re.match(r"wd_s(\d+)", ins["objectId"])
        if m:
            texts.setdefault(int(m.group(1)), []).append(ins["text"])
    n = max(texts, default=0)
    return [re.sub(r"\s+", " ", " ".join(texts.get(i, []))).strip() for i in range(1, n + 1)]


def load_requests(path: Path) -> list[dict[str, Any]]:
    data = json.loads(Path(path).read_text())
    return data["requests"] if isinstance(data, dict) else data


def chunk_problems(full: Path) -> list[str]:
    """fills.json and fills-NN.json must agree, and each chunk must fit one connector call."""
    chunks = sorted(full.parent.glob(full.stem + "-[0-9][0-9].json"))
    if not chunks:
        return [f"no {full.stem}-NN.json chunks next to {full.name}"]
    joined: list[Any] = []
    problems = []
    for c in chunks:
        reqs = load_requests(c)
        if len(reqs) > MAX_CHUNK:
            problems.append(f"{c.name} has {len(reqs)} requests (over {MAX_CHUNK})")
        joined += reqs
    if joined != load_requests(full):
        problems.append("chunks do not add up to the full fills batch")
    return problems


# ---------- PowerPoint ----------

def pptx_problems(path: Path, template: Path = ROOT / "assets" / "template.pptx") -> dict[str, list[str]]:
    from pptx import Presentation
    from pptx.enum.dml import MSO_COLOR_TYPE

    prs = Presentation(str(path))
    layouts = {l.name.strip() for l in Presentation(str(template)).slide_layouts}
    out: dict[str, list[str]] = {"layouts": [], "fonts": [], "colors": [], "footer": []}

    def color(c, where):
        try:
            if c.type == MSO_COLOR_TYPE.RGB and str(c.rgb).upper() not in BRAND_HEX:
                out["colors"].append(f"{where}: #{c.rgb}")
        except (AttributeError, TypeError):
            pass

    def text_frame(tf, where):
        for p in tf.paragraphs:
            for r in p.runs:
                if r.font.name and r.font.name not in PPTX_FONTS:
                    out["fonts"].append(f"{where}: {r.font.name}")
                color(r.font.color, where)

    for i, slide in enumerate(prs.slides, 1):
        lname = slide.slide_layout.name.strip()
        if lname not in layouts:
            out["layouts"].append(f"slide {i}: layout {lname!r} is not in the template")
        texts = []
        for shape in slide.shapes:
            where = f"slide {i} {shape.name}"
            if shape.has_text_frame:
                text_frame(shape.text_frame, where)
                texts.append(shape.text_frame.text)
            try:
                if shape.fill.type is not None:
                    color(shape.fill.fore_color, where)
            except (AttributeError, TypeError, NotImplementedError):
                pass
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        text_frame(cell.text_frame, where + " cell")
                        try:
                            if cell.fill.type is not None:
                                color(cell.fill.fore_color, where + " cell fill")
                        except (AttributeError, TypeError, NotImplementedError):
                            pass
            if getattr(shape, "has_chart", False) and shape.has_chart:
                for plot in shape.chart.plots:
                    for s in plot.series:
                        try:
                            color(s.format.fill.fore_color, where + f" series {s.name}")
                        except (AttributeError, TypeError, NotImplementedError):
                            pass
        if lname not in NO_FOOTER_LAYOUTS and not any("Workday Confidential" in t for t in texts):
            out["footer"].append(f"slide {i}: no 'Workday Confidential' footer")
    return {k: sorted(set(v))[:10] for k, v in out.items()}


def pptx_text_by_slide(path: Path) -> list[str]:
    from pptx import Presentation

    texts = []
    for slide in Presentation(str(path)).slides:
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                parts.extend(c.text for row in shape.table.rows for c in row.cells)
            if getattr(shape, "has_chart", False) and shape.has_chart:
                plot = shape.chart.plots[0]
                parts.extend(str(c) for c in plot.categories)
                for s in plot.series:
                    parts.append(str(s.name))
                    parts.extend(f"{v:g}" for v in s.values if v is not None)
        if slide.has_notes_slide:
            parts.append(slide.notes_slide.notes_text_frame.text)
        texts.append(re.sub(r"\s+", " ", " ".join(parts)).strip())
    return texts


def pptx_layouts(path: Path) -> list[str]:
    from pptx import Presentation

    return [s.slide_layout.name.strip() for s in Presentation(str(path)).slides]
