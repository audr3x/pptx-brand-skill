from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from brand import BRAND
from build_deck import build
from build_slides import bind_fills, build_plan
from validate_brand import validate_brand


def all_patterns(image: str) -> dict:
    people = [
        {
            "name": f"Speaker {index}",
            "role": "Role",
            "company": "Workday",
            "image": image,
        }
        for index in range(1, 5)
    ]
    items = [
        {"title": f"Item {index}", "body": f"Body {index}"}
        for index in range(1, 7)
    ]
    return {
        "title": "All Workday layouts",
        "footer": "Workday Confidential",
        "slides": [
            {"layout": "title", "title": "Title", "subtitle": "Subtitle"},
            {"layout": "intro-one", "speakers": people[:1]},
            {"layout": "intro-three", "speakers": people[:3]},
            {"layout": "intro-four", "speakers": people},
            {"layout": "content", "title": "Content", "bullets": ["One", "Two"]},
            {"layout": "title-only", "title": "Title only", "body": "Body"},
            {
                "layout": "title-only-alt",
                "title": "Warm title",
                "subtitle": "Subhead",
                "body": "Body",
            },
            {
                "layout": "feature-grid",
                "title": "Feature grid",
                "subtitle": "Subhead",
                "items": items[:3],
            },
            {"layout": "two-column", "title": "Three items", "items": items[:3]},
            {
                "layout": "two-column-alt",
                "title": "Six items",
                "items": items,
            },
            {"layout": "section", "title": "Section", "eyebrow": "Eyebrow"},
            {"layout": "agenda", "title": "Agenda", "items": items},
            {"layout": "rail", "title": "Rail", "body": "Body"},
            {
                "layout": "three-photos",
                "title": "Photos",
                "images": [image, image, image],
            },
            {
                "layout": "photo-right",
                "title": "Photo right",
                "image": image,
                "items": items[:3],
            },
            {"layout": "closing", "title": "Thank you"},
            {"layout": "blank", "title": "Blank", "body": "Body"},
            {
                "layout": "table",
                "title": "Table",
                "headers": ["A", "B"],
                "rows": [["1", "2"]],
            },
            {
                "layout": "chart",
                "title": "Chart",
                "chart": {
                    "type": "column",
                    "categories": ["A", "B"],
                    "series": [{"name": "Series", "values": [1, 2]}],
                },
            },
            {"layout": "timeline", "title": "Timeline", "items": items[:4]},
        ],
    }


class BrandSystemTests(unittest.TestCase):
    def test_registry_matches_template_and_skills(self):
        errors = validate_brand(ROOT / "assets" / "template.pptx")
        self.assertEqual([], errors)

    def test_registry_matches_original_when_available(self):
        source = Path(
            "/Users/audrey/Downloads/Workday Corporate PowerPoint Template 2.pptx"
        )
        if not source.exists():
            self.skipTest("Original Workday deck is not available")
        errors = validate_brand(ROOT / "assets" / "template.pptx", source)
        self.assertEqual([], errors)

    def test_canonical_blue_sky_uses_theme_hex(self):
        self.assertEqual("9ECFFF", BRAND.colors["blue_sky"])
        with self.assertRaises(ValueError):
            BRAND.color_hex("9FCFFF")

    def test_outline_limits_fail_at_the_boundary(self):
        with self.assertRaisesRegex(ValueError, "exactly 3 speakers"):
            BRAND.normalize_slide(
                {"layout": "intro-three", "speakers": [{"name": "Only one"}]}
            )
        with self.assertRaisesRegex(ValueError, "at most 3 items"):
            BRAND.normalize_slide(
                {
                    "layout": "two-column",
                    "title": "Too many",
                    "items": ["1", "2", "3", "4"],
                }
            )
        with self.assertRaisesRegex(ValueError, "exactly 3 images"):
            BRAND.normalize_slide(
                {
                    "layout": "three-photos",
                    "title": "Too few",
                    "images": ["one.png"],
                }
            )
        legacy = BRAND.normalize_slide(
            {
                "layout": "intro",
                "title": "Legacy name",
                "subtitle": "Legacy role",
                "org": "Legacy company",
            }
        )
        self.assertEqual("Legacy name", legacy.get("speakers")[0]["name"])

    def test_powerpoint_builds_every_pattern(self):
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            image = directory_path / "image.png"
            Image.new("RGB", (80, 80), "#1C98E8").save(image)
            outline = all_patterns(str(image))
            output = directory_path / "all-layouts.pptx"
            build(outline, ROOT / "assets" / "template.pptx", output)
            presentation = Presentation(str(output))
            normalized = BRAND.normalize_outline(outline)
            self.assertEqual(len(normalized), len(presentation.slides))
            self.assertEqual(
                [slide.pattern.layout.name for slide in normalized],
                [slide.slide_layout.name for slide in presentation.slides],
            )
            text = "\n".join(
                shape.text
                for slide in presentation.slides
                for shape in slide.shapes
                if getattr(shape, "has_text_frame", False)
            )
            for speaker in ("Speaker 1", "Speaker 2", "Speaker 3", "Speaker 4"):
                self.assertIn(speaker, text)
            self.assertNotIn("Click to edit", text)
            self.assertNotIn("Insert slide title", text)
            footer_count = sum(
                shape.text == "Workday Confidential"
                for slide in presentation.slides
                for shape in slide.shapes
                if getattr(shape, "has_text_frame", False)
            )
            self.assertEqual(
                sum(slide.pattern.layout.footer for slide in normalized),
                footer_count,
            )
            for slide in presentation.slides:
                for shape in slide.shapes:
                    self.assertGreaterEqual(shape.left, 0)
                    self.assertGreaterEqual(shape.top, 0)
                    self.assertLessEqual(shape.left + shape.width, presentation.slide_width)
                    self.assertLessEqual(shape.top + shape.height, presentation.slide_height)

    def test_slides_plan_covers_every_pattern(self):
        outline = all_patterns("https://example.com/image.png")
        plan = build_plan(outline)
        self.assertEqual(len(outline["slides"]), len(plan["slides"]))
        self.assertEqual(
            [slide.pattern.layout.name for slide in BRAND.normalize_outline(outline)],
            [slide["layout"] for slide in plan["slides"]],
        )
        requests = json.dumps(plan["requests"])
        self.assertIn('"createImage"', requests)
        self.assertIn('"wd_s3_speaker_2_company"', requests)
        self.assertIn('"wd_s4_speaker_3_company"', requests)

    def test_live_request_binding_preserves_connector_ids(self):
        outline = {
            "slides": [
                {"layout": "content", "title": "Title", "bullets": ["Body"]}
            ]
        }
        planned = build_plan(outline)["slidePlans"][0]
        presentation = {
            "slides": [
                {
                    "objectId": "live-slide",
                    "pageElements": [
                        {
                            "objectId": "live-title",
                            "shape": {
                                "placeholder": {"type": "TITLE", "index": 0}
                            },
                        },
                        {
                            "objectId": "live-body",
                            "shape": {
                                "placeholder": {"type": "BODY", "index": 0}
                            },
                        },
                    ],
                }
            ]
        }
        rewritten = bind_fills(
            {"slidePlans": [planned]},
            presentation,
        )
        payload = json.dumps(rewritten)
        self.assertIn("live-title", payload)
        self.assertIn("live-body", payload)
        self.assertIn("live-slide", payload)


if __name__ == "__main__":
    unittest.main()
