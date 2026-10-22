"""python3 -m unittest discover -s tests"""
from __future__ import annotations

import sys
import unittest
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import build_slides as bs  # noqa: E402

PUBLIC_RAW = "https://raw.githubusercontent.com/acme/workday-brand-assets/main/cover-gradient.png"
PRIVATE_RAW = "https://raw.githubusercontent.com/acme/private-repo/main/cover-gradient.png"
PUBLIC_SMILE = "https://cdn.example.com/workday-smile.png"
OUTLINE = bs.load_outline(ROOT / "examples" / "outline.yaml")


def fake_probe(url: str, timeout: float = 6.0):
    if "private-repo" in url:
        return 404, "text/plain"
    if "signin" in url:
        return 200, "text/html"
    return 200, "image/png"


def args(**kw):
    base = dict(
        cover_image_id=None, smile_image_id=None, cover_image_url=None, smile_image_url=None,
        assets_ids=None, skip_url_check=False, place_smile=False,
    )
    base.update(kw)
    return mock.Mock(**base)


def image_urls(chrome):
    plan = bs.build_plan(OUTLINE, chrome)
    return {r["createImage"]["objectId"]: r["createImage"]["url"] for r in plan["requests"] if "createImage" in r}, plan


def title_color(plan):
    for r in plan["requests"]:
        style = r.get("updateTextStyle")
        if style and style["objectId"] == "wd_s1_title":
            return style["style"]["foregroundColor"]["opaqueColor"]["rgbColor"]
    raise AssertionError("no title style")


INK = bs.rgb_frac("ink")
PAPER = bs.rgb_frac("paper")


@mock.patch.dict("os.environ", {}, clear=True)
@mock.patch.object(bs, "probe_url", side_effect=fake_probe)
class ImageUrlTests(unittest.TestCase):
    def test_public_urls_are_used_as_given(self, _):
        chrome = bs.resolve_chrome(args(cover_image_url=PUBLIC_RAW, smile_image_url=PUBLIC_SMILE, place_smile=True))
        urls, plan = image_urls(chrome)
        self.assertEqual(urls, {"wd_s1_cover": PUBLIC_RAW, "wd_s1_smile": PUBLIC_SMILE, "wd_s7_smile": PUBLIC_SMILE})
        self.assertEqual(title_color(plan), INK)

    def test_smile_skipped_by_default_because_layouts_carry_it(self, _):
        urls, _ = image_urls(bs.resolve_chrome(args(cover_image_url=PUBLIC_RAW, smile_image_url=PUBLIC_SMILE, smile_image_id="S1")))
        self.assertEqual(urls, {"wd_s1_cover": PUBLIC_RAW})

    def test_placeholder_ids_never_reach_create_image(self, probe):
        urls, plan = image_urls(bs.resolve_chrome(args(cover_image_id="YOUR_COVER_IMAGE_ID", cover_image_url="<cover URL>")))
        self.assertEqual(urls, {})
        self.assertEqual(title_color(plan), PAPER)
        probe.assert_not_called()

    def test_public_url_wins_over_drive_id(self, _):
        urls, _ = image_urls(bs.resolve_chrome(args(cover_image_url=PUBLIC_RAW, cover_image_id="DRIVE1")))
        self.assertEqual(urls["wd_s1_cover"], PUBLIC_RAW)

    def test_private_raw_falls_back_to_drive_id(self, _):
        urls, _ = image_urls(bs.resolve_chrome(args(cover_image_url=PRIVATE_RAW, cover_image_id="DRIVE1")))
        self.assertEqual(urls["wd_s1_cover"], bs.drive_download_url("DRIVE1"))

    def test_private_raw_without_id_gives_flat_cover(self, _):
        urls, plan = image_urls(bs.resolve_chrome(args(cover_image_url=PRIVATE_RAW)))
        self.assertEqual(urls, {})
        self.assertEqual(title_color(plan), PAPER)

    def test_no_images_is_ballpoint_with_paper_type(self, probe):
        urls, plan = image_urls(bs.resolve_chrome(args()))
        self.assertEqual(urls, {})
        probe.assert_not_called()
        fills = [r for r in plan["requests"] if "updatePageProperties" in r]
        colors = {r["updatePageProperties"]["objectId"]: r["updatePageProperties"]["pageProperties"]["pageBackgroundFill"]["solidFill"]["color"]["rgbColor"] for r in fills}
        self.assertEqual(colors["wd_s1"], bs.rgb_frac("ballpoint"))
        self.assertEqual(colors["wd_s7"], bs.rgb_frac("ink"))
        self.assertEqual(title_color(plan), PAPER)

    def test_failing_non_github_url_falls_back_to_drive_id(self, _):
        bad = "https://intranet.example.com/signin/cover.png"
        urls, _ = image_urls(bs.resolve_chrome(args(cover_image_url=bad, cover_image_id="C1")))
        self.assertEqual(urls["wd_s1_cover"], bs.drive_download_url("C1"))
        urls, _ = image_urls(bs.resolve_chrome(args(cover_image_url=bad)))
        self.assertEqual(urls["wd_s1_cover"], bad)

    def test_skip_url_check_keeps_url_without_fetching(self, probe):
        urls, _ = image_urls(bs.resolve_chrome(args(cover_image_url=PRIVATE_RAW, skip_url_check=True)))
        self.assertEqual(urls["wd_s1_cover"], PRIVATE_RAW)
        probe.assert_not_called()


def requests_for(object_prefix: str):
    plan = bs.build_plan(OUTLINE, {})
    return [r for r in plan["requests"] if object_prefix in json_ids(r)]


def json_ids(r):
    v = next(iter(r.values()))
    ids = [v.get("objectId", "")] + list(v.get("pageElementObjectIds", []))
    return " ".join(i for i in ids if i)


class LiveFindingTests(unittest.TestCase):
    def test_cover_title_is_bold_700(self):
        style = next(r["updateTextStyle"]["style"] for r in requests_for("wd_s1_title") if "updateTextStyle" in r)
        self.assertEqual(style["weightedFontFamily"]["weight"], 700)

    def test_bullets_have_spacing_that_lists_do_not_collapse(self):
        reqs = requests_for("wd_s3_body")
        kinds = [next(iter(r)) for r in reqs]
        self.assertLess(kinds.index("createParagraphBullets"), kinds.index("updateParagraphStyle"))
        para = next(r["updateParagraphStyle"] for r in reqs if "updateParagraphStyle" in r)
        self.assertEqual(para["style"]["spacingMode"], "NEVER_COLLAPSE")
        self.assertEqual(para["style"]["spaceAbove"]["magnitude"], 6)
        self.assertEqual(para["style"]["spaceBelow"]["magnitude"], 6)
        self.assertEqual(set(para["fields"].split(",")), {"spaceAbove", "spaceBelow", "spacingMode"})
        sizes = {r["updateTextStyle"]["style"]["fontSize"]["magnitude"] for r in reqs if "updateTextStyle" in r}
        self.assertEqual(sizes, {20})

    def test_chart_has_axis_ticks_and_gridlines(self):
        reqs = requests_for("wd_s6_")
        ticks = [r["insertText"]["text"] for r in reqs if "insertText" in r and "_tick_" in r["insertText"]["objectId"]]
        self.assertEqual(ticks, ["0", "5", "10", "15", "20", "25"])
        self.assertTrue(any(r.get("createShape", {}).get("objectId") == "wd_s6_yaxis" for r in reqs))
        grid_first = min(i for i, r in enumerate(reqs) if r.get("createShape", {}).get("objectId", "").startswith("wd_s6_grid_"))
        bar_first = min(i for i, r in enumerate(reqs) if r.get("createShape", {}).get("objectId", "").startswith("wd_s6_bar_"))
        self.assertLess(grid_first, bar_first)

    def test_axis_scale(self):
        self.assertEqual(bs.axis_scale(0, 22), (5, 0, 5))
        self.assertEqual(bs.axis_scale(0, 100), (20, 0, 5))
        self.assertEqual(bs.axis_scale(0, 0.8), (0.2, 0, 4))
        self.assertEqual(bs.axis_scale(-20, -10), (5, -4, 0))
        self.assertEqual(bs.axis_scale(-3, 10), (5, -1, 2))

    def test_negative_bars_hang_below_zero(self):
        outline = {"slides": [{"layout": "chart", "title": "Margin",
                               "chart": {"categories": ["A", "B"], "series": [{"name": "M", "values": [10, -5]}]}}]}
        reqs = bs.build_plan(outline, {})["requests"]
        shapes = {r["createShape"]["objectId"]: r["createShape"]["elementProperties"] for r in reqs if "createShape" in r}
        texts = {r["insertText"]["objectId"]: r["insertText"]["text"] for r in reqs if "insertText" in r}
        zero_tick = next(k for k, v in texts.items() if "_tick_" in k and v == "0")
        zero_y = shapes[zero_tick]["transform"]["translateY"] / bs.EMU_PER_INCH + 0.14
        def box(oid):
            e = shapes[oid]
            y = e["transform"]["translateY"] / bs.EMU_PER_INCH
            return y, y + e["size"]["height"]["magnitude"] / bs.EMU_PER_INCH
        pos_top, pos_bottom = box("wd_s1_bar_0_0")
        neg_top, neg_bottom = box("wd_s1_bar_1_0")
        self.assertAlmostEqual(pos_bottom, zero_y, places=3)
        self.assertAlmostEqual(neg_top, zero_y, places=3)
        self.assertAlmostEqual((neg_bottom - neg_top) * 2, pos_bottom - pos_top, places=3)

    def test_empty_footer_means_no_footer(self):
        outline = {"slides": [{"layout": "content", "title": "T", "bullets": ["a"], "footer": ""}]}
        reqs = bs.build_plan(outline, {})["requests"]
        self.assertFalse(any("_footer" in json_ids(r) for r in reqs))

    def test_section_footer_clears_rail(self):
        footer = next(r["createShape"] for r in requests_for("wd_s2_footer") if "createShape" in r)
        x = footer["elementProperties"]["transform"]["translateX"] / bs.EMU_PER_INCH
        self.assertGreater(x, bs.SECTION_RAIL_WIDTH_IN)

    def test_template_section_footer_placeholder_clears_rail(self):
        from pptx import Presentation
        from pptx.enum.shapes import PP_PLACEHOLDER

        prs = Presentation(str(ROOT / "assets" / "template.pptx"))
        layout = next(l for l in prs.slide_layouts if l.name.strip() == "Section Title")
        footer = next(p for p in layout.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.FOOTER)
        self.assertGreater(footer.left / bs.EMU_PER_INCH, bs.SECTION_RAIL_WIDTH_IN)


if __name__ == "__main__":
    unittest.main()
