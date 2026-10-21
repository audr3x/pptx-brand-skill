#!/usr/bin/env python3
"""Build an on-brand Workday deck from a JSON or YAML outline.

Usage:
    python3 scripts/build_deck.py examples/outline.yaml -o examples/example-deck.pptx
    python3 scripts/build_deck.py outline.json --template assets/template.pptx
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Any, Mapping

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.table import Table
from pptx.util import Inches, Pt

from brand import BRAND, Composition, NormalizedSlide, load_outline

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = ROOT / "assets" / "template.pptx"

MAX_BULLETS = BRAND.limits["bullets"]
CORE = BRAND.core
SECONDARY = BRAND.secondary
CHART_SERIES = BRAND.chart_series
FONT_TITLE = BRAND.fonts["title"]
FONT_BODY = BRAND.fonts["body"]


def rgb(name_or_hex: str) -> RGBColor:
    return RGBColor.from_string(BRAND.color_hex(name_or_hex))


def find_layout(prs: Presentation, name: str):
    wanted = name.strip().lower()
    for layout in prs.slide_layouts:
        if layout.name.strip().lower() == wanted:
            return layout
    available = [f"{i}:{ly.name!r}" for i, ly in enumerate(prs.slide_layouts)]
    raise SystemExit(f"Unknown layout {name!r}. Available: {available}")


def iter_placeholders(slide):
    for shape in slide.shapes:
        if shape.is_placeholder:
            yield shape


def placeholder_of_types(slide, types):
    for shape in iter_placeholders(slide):
        if shape.placeholder_format.type in types:
            return shape
    return None


def set_run_font(run, *, name=FONT_TITLE, size=None, bold=None, color=None):
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = rgb(color)


def replace_text(shape, text: str | None, *, keep_format=True):
    """Replace placeholder text. Empty string clears the prompt."""
    value = "" if text is None else str(text)
    tf = shape.text_frame
    if keep_format and tf.paragraphs and tf.paragraphs[0].runs:
        tf.paragraphs[0].runs[0].text = value
        for run in tf.paragraphs[0].runs[1:]:
            run.text = ""
        # Drop leftover paragraphs by emptying them.
        for para in tf.paragraphs[1:]:
            if para.runs:
                para.runs[0].text = ""
                for run in para.runs[1:]:
                    run.text = ""
            else:
                para.text = ""
        return
    shape.text = value


def add_textbox(
    slide,
    left,
    top,
    width,
    height,
    text,
    *,
    font=FONT_TITLE,
    size=16,
    bold=False,
    color="ink",
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run_font(run, name=font, size=size, bold=bold, color=color)
    return box


def fill_footer(slide, text: str | None, layout):
    """python-pptx does not clone footer placeholders. Draw the brand footer ourselves."""
    if text is None or not layout.footer:
        return
    color = "ink" if layout.name in {"Section Title", "Layout 1"} else "staple"
    footer = placeholder_of_types(slide, {PP_PLACEHOLDER.FOOTER})
    if footer is not None:
        replace_text(footer, text)
        if footer.has_text_frame:
            for para in footer.text_frame.paragraphs:
                for run in para.runs:
                    set_run_font(run, name=FONT_TITLE, size=8, bold=False, color=color)
        return
    left, top, width, height = layout.footer_region
    add_textbox(
        slide,
        left,
        top,
        width,
        height,
        text,
        font=FONT_TITLE,
        size=8,
        bold=False,
        color=color,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def fill_title(slide, text: str | None):
    title = placeholder_of_types(
        slide,
        {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE},
    )
    if title is not None and text is not None:
        replace_text(title, text)


def fill_subtitle(slide, text: str | None):
    sub = placeholder_of_types(slide, {PP_PLACEHOLDER.SUBTITLE})
    if sub is not None:
        replace_text(sub, text or "")


def body_placeholder(slide):
    for shape in iter_placeholders(slide):
        kind = shape.placeholder_format.type
        if kind in {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT} and shape.has_text_frame:
            # Skip footer / title-like bodies already handled.
            if kind == PP_PLACEHOLDER.FOOTER:
                continue
            return shape
    return None


def placeholder_by_idx(slide, idx: int):
    for shape in iter_placeholders(slide):
        if shape.placeholder_format.idx == idx:
            return shape
    return None


def set_bullets(shape, items: list[Any]):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True

    def walk(entries, level):
        rows = []
        for entry in entries:
            if isinstance(entry, str):
                rows.append((entry, level))
            elif isinstance(entry, dict):
                text = entry.get("text") or entry.get("title") or ""
                rows.append((text, level))
                kids = entry.get("children") or entry.get("bullets") or []
                rows.extend(walk(kids, min(level + 1, 2)))
            else:
                rows.append((str(entry), level))
        return rows

    rows = walk(items, 0)
    if len(rows) > MAX_BULLETS:
        print(
            f"warning: {len(rows)} bullets on one slide; brand limit is about {MAX_BULLETS}. Truncating.",
            file=sys.stderr,
        )
        rows = rows[:MAX_BULLETS]
    if not rows:
        rows = [("", 0)]
    for i, (text, level) in enumerate(rows):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = level
        p.alignment = PP_ALIGN.LEFT
        # One run so the layout bullet style can attach.
        if p.runs:
            p.runs[0].text = text
            set_run_font(
                p.runs[0],
                name=FONT_BODY,
                size=20 if level == 0 else 16,
                bold=False,
                color="ink",
            )
        else:
            run = p.add_run()
            run.text = text
            set_run_font(
                run,
                name=FONT_BODY,
                size=20 if level == 0 else 16,
                bold=False,
                color="ink",
            )


def image_path(image: Any) -> Path:
    if isinstance(image, dict):
        image = image.get("path") or image.get("file")
    if not image:
        raise ValueError("A local image path is required for PowerPoint output")
    path = Path(image)
    if not path.is_file():
        path = (ROOT / image).resolve()
    if not path.is_file():
        raise ValueError(f"Image not found: {image}")
    return path


def fill_picture_slots(slide, indexes: list[int] | tuple[int, ...], images: list[Any]):
    for idx, image in zip(indexes, images):
        placeholder = placeholder_by_idx(slide, idx)
        if placeholder is None:
            raise ValueError(f"Picture placeholder {idx} is missing")
        placeholder.insert_picture(str(image_path(image)))


def _set_cell_border(cell, *, color_hex: str, width_emu: int, edges=None):
    tcPr = cell._tc.get_or_add_tcPr()
    for edge in edges or ("lnL", "lnR", "lnT", "lnB"):
        ln = tcPr.find(qn(f"a:{edge}"))
        if ln is None:
            ln = tcPr.makeelement(qn(f"a:{edge}"), {})
            tcPr.append(ln)
        ln.set("w", str(width_emu))
        for child in list(ln):
            ln.remove(child)
        sf = ln.makeelement(qn("a:solidFill"), {})
        sf.append(sf.makeelement(qn("a:srgbClr"), {"val": color_hex}))
        ln.append(sf)


def style_table(table: Table, headers: list[str], rows: list[list[Any]]):
    def paint(cell, *, fill, text, font, size, bold, color):
        cell.text = ""
        tf = cell.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.10)
        tf.margin_right = Inches(0.08)
        tf.margin_top = Inches(0.08)
        tf.margin_bottom = Inches(0.08)
        try:
            tf._txBody.bodyPr.set("anchor", "ctr")
        except Exception:
            pass
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = "" if text is None else str(text)
        set_run_font(run, name=font, size=size, bold=bold, color=color)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(fill)

    n_cols = len(headers)
    last_row = len(rows)
    for col, header in enumerate(headers):
        cell = table.cell(0, col)
        paint(
            cell,
            fill="water_cooler",
            text=header,
            font=FONT_TITLE,
            size=12,
            bold=True,
            color="paper",
        )
        # Header reads as one bar. Hairline Paper between cells, Desk rule under.
        _set_cell_border(cell, color_hex=CORE["water_cooler"], width_emu=6350)
        _set_cell_border(cell, color_hex=SECONDARY["desk"], width_emu=19050, edges=("lnB",))
        if col < n_cols - 1:
            _set_cell_border(cell, color_hex=CORE["paper"], width_emu=6350, edges=("lnR",))
    for r, row in enumerate(rows, start=1):
        # Keyboard (cream) on odd body rows so they do not vanish into Paper.
        fill = "keyboard" if r % 2 else "business_card"
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            paint(
                cell,
                fill=fill,
                text=value,
                font=FONT_BODY,
                size=11,
                bold=False,
                color="ink",
            )
            _set_cell_border(cell, color_hex=SECONDARY["desk"], width_emu=6350)
            if r == last_row:
                _set_cell_border(cell, color_hex=SECONDARY["desk"], width_emu=12700, edges=("lnB",))


def add_table_slide(slide, spec: dict[str, Any]):
    headers = [str(h) for h in spec.get("headers") or []]
    rows = spec.get("rows") or []
    if not headers:
        raise SystemExit("table slides need a 'headers' list")
    n_rows = 1 + len(rows)
    n_cols = len(headers)
    normalized = []
    for row in rows:
        values = [row.get(h, "") if isinstance(row, dict) else (row[i] if i < len(row) else "") for i, h in enumerate(headers)]
        if isinstance(row, dict):
            values = [row.get(h, "") for h in headers]
        else:
            values = list(row) + [""] * (n_cols - len(row))
            values = values[:n_cols]
        normalized.append(values)
    # Keep rows even. Short tables stay compact rather than stretching the last row.
    height = max(1.70, min(3.40, 0.48 * n_rows))
    table_shape = slide.shapes.add_table(
        n_rows,
        n_cols,
        Inches(0.50),
        Inches(1.34),
        Inches(9.00),
        Inches(height),
    )
    table = table_shape.table
    row_h = int(table_shape.height / n_rows)
    for row in table.rows:
        row.height = row_h
    style_table(table, headers, normalized)


def add_chart_slide(slide, spec: dict[str, Any]):
    chart_spec = spec.get("chart") or {}
    categories = chart_spec.get("categories") or []
    series_list = chart_spec.get("series") or []
    if not categories or not series_list:
        raise SystemExit("chart slides need chart.categories and chart.series")
    kind_name = (chart_spec.get("type") or "column").lower()
    kinds = {
        "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "bar": XL_CHART_TYPE.BAR_CLUSTERED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "pie": XL_CHART_TYPE.PIE,
    }
    if kind_name not in kinds:
        raise SystemExit(f"Unsupported chart type {kind_name!r}. Use column, bar, line, or pie.")
    data = CategoryChartData()
    data.categories = categories
    for series in series_list:
        data.add_series(series.get("name") or "Series", series.get("values") or [])
    chart = slide.shapes.add_chart(
        kinds[kind_name],
        Inches(0.50),
        Inches(1.28),
        Inches(9.00),
        Inches(3.70),
        data,
    ).chart
    chart.has_legend = len(series_list) > 1
    if chart.has_legend:
        chart.legend.include_in_layout = True
        chart.legend.position = XL_LEGEND_POSITION.RIGHT
        chart.legend.font.size = Pt(10)
        chart.legend.font.color.rgb = rgb("ink")
        chart.legend.font.name = FONT_TITLE
    plot = chart.plots[0]
    plot.has_data_labels = False
    for i, series in enumerate(chart.series):
        color = CHART_SERIES[i % len(CHART_SERIES)]
        try:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = RGBColor.from_string(color)
        except Exception:
            pass
        try:
            series.format.line.color.rgb = RGBColor.from_string(color)
        except Exception:
            pass
    def _axis_font(axis):
        try:
            axis.tick_labels.font.size = Pt(10)
            axis.tick_labels.font.color.rgb = rgb("ink")
            axis.tick_labels.font.name = FONT_BODY
        except Exception:
            pass
        try:
            axis.format.line.color.rgb = rgb("desk")
        except Exception:
            pass
    try:
        _axis_font(chart.category_axis)
        _axis_font(chart.value_axis)
    except Exception:
        pass
    try:
        chart.value_axis.has_major_gridlines = True
        chart.value_axis.major_gridlines.format.line.color.rgb = rgb("desk")
    except Exception:
        pass
    # Plot / wall stay Paper when the format API allows it.
    try:
        chart.plot_area.format.fill.solid()
        chart.plot_area.format.fill.fore_color.rgb = rgb("paper")
    except Exception:
        pass


def add_section_copy(slide, spec: dict[str, Any]):
    eyebrow = spec.get("eyebrow") or spec.get("kicker")
    title = spec.get("title") or spec.get("statement") or ""
    if eyebrow:
        left, top, width, height = BRAND.region("section_eyebrow")
        add_textbox(
            slide,
            left,
            top,
            width,
            height,
            str(eyebrow).upper(),
            font=FONT_TITLE,
            size=14,
            bold=False,
            color="water_cooler",
        )
    left, top, width, height = BRAND.region("section_title")
    add_textbox(
        slide,
        left,
        top,
        width,
        height,
        title,
        font=FONT_TITLE,
        size=28,
        bold=True,
        color="ink",
        anchor=MSO_ANCHOR.TOP,
    )


def item_copy(item: Any) -> tuple[str, str]:
    if isinstance(item, str):
        return item, ""
    return (
        item.get("title") or item.get("heading") or "",
        item.get("body") or item.get("text") or "",
    )


def add_two_column_items(slide, items: list[Any], *, compact: bool):
    stack = "two_column_compact" if compact else "two_column_numbered"
    for i, item in enumerate(items):
        heading, body = item_copy(item)
        if not compact:
            left, top, width, height = BRAND.stack_region(stack, "number", i)
            add_textbox(
                slide,
                left,
                top,
                width,
                height,
                f"{i + 1:02d}",
                font=FONT_TITLE,
                size=20,
                bold=True,
                color="water_cooler",
            )
        left, top, width, height = BRAND.stack_region(stack, "heading", i)
        add_textbox(
            slide,
            left,
            top,
            width,
            height,
            heading,
            font=FONT_TITLE,
            size=13 if compact else 15,
            bold=True,
            color="ink",
        )
        if body:
            left, top, width, height = BRAND.stack_region(stack, "body", i)
            add_textbox(
                slide,
                left,
                top,
                width,
                height,
                body,
                font=FONT_BODY,
                size=11 if compact else 13,
                bold=False,
                color="ink",
            )


def add_cover_date(slide, date_text: str | None):
    if not date_text:
        return
    left, top, width, height = BRAND.region("cover_date")
    add_textbox(
        slide,
        left,
        top,
        width,
        height,
        date_text,
        font=FONT_TITLE,
        size=12,
        bold=False,
        color="ink",
    )


def add_closing_caption(slide, text: str | None):
    if not text:
        return
    left, top, width, height = BRAND.region("closing_caption")
    add_textbox(
        slide,
        left,
        top,
        width,
        height,
        text,
        font=FONT_TITLE,
        size=16,
        bold=False,
        color="paper",
        align=PP_ALIGN.CENTER,
    )


def fill_intro(slide, normalized: NormalizedSlide):
    for slot, speaker in zip(
        normalized.pattern.layout.speaker_slots,
        normalized.get("speakers") or [],
    ):
        for field in ("name", "role", "company"):
            placeholder = placeholder_by_idx(slide, slot[field])
            if placeholder is None:
                raise ValueError(f"Speaker placeholder {slot[field]} is missing")
            replace_text(placeholder, speaker.get(field) or "")
        if speaker.get("image"):
            fill_picture_slots(slide, [slot["picture"]], [speaker["image"]])


def add_feature_grid(slide, spec: Mapping[str, Any]):
    subtitle = spec.get("subtitle")
    if subtitle:
        left, top, width, height = BRAND.region("title_alt_subhead")
        add_textbox(
            slide,
            left,
            top,
            width,
            height,
            str(subtitle),
            font=FONT_TITLE,
            size=16,
            color="ink",
        )
    for i, item in enumerate(spec.get("items") or []):
        heading, body = item_copy(item)
        left, top, width, height = BRAND.stack_region("feature_grid", "heading", i)
        add_textbox(
            slide,
            left,
            top,
            width,
            height,
            heading,
            font=FONT_TITLE,
            size=15,
            bold=True,
            color="ink",
        )
        if body:
            left, top, width, height = BRAND.stack_region("feature_grid", "body", i)
            add_textbox(
                slide,
                left,
                top,
                width,
                height,
                body,
                font=FONT_BODY,
                size=11,
                color="ink",
            )


def add_photo_right_items(slide, items: list[Any]):
    for i, item in enumerate(items):
        heading, body = item_copy(item)
        left, top, width, height = BRAND.stack_region(
            "photo_right_items", "heading", i
        )
        add_textbox(
            slide,
            left,
            top,
            width,
            height,
            heading,
            font=FONT_TITLE,
            size=13,
            bold=True,
            color="ink",
        )
        if body:
            left, top, width, height = BRAND.stack_region(
                "photo_right_items", "body", i
            )
            add_textbox(
                slide,
                left,
                top,
                width,
                height,
                body,
                font=FONT_BODY,
                size=11,
                color="ink",
            )


def add_agenda(slide, spec: Mapping[str, Any]):
    left, top, width, height = BRAND.region("section_eyebrow")
    add_textbox(
        slide,
        left,
        top,
        width,
        height,
        str(spec.get("title") or "").upper(),
        font=FONT_TITLE,
        size=14,
        color="water_cooler",
    )
    for i, item in enumerate(spec.get("items") or []):
        heading, body = item_copy(item)
        number = str(item.get("number") if isinstance(item, dict) else i + 1)
        left, top, width, height = BRAND.stack_region(
            "agenda_items", "number", i
        )
        add_textbox(
            slide,
            left,
            top,
            width,
            height,
            number.zfill(2),
            font=FONT_TITLE,
            size=13,
            bold=True,
            color="water_cooler",
        )
        left, top, width, height = BRAND.stack_region(
            "agenda_items", "heading", i
        )
        text = heading if not body else f"{heading}\n{body}"
        add_textbox(
            slide,
            left,
            top,
            width,
            height + (0.2 if body else 0),
            text,
            font=FONT_TITLE if not body else FONT_BODY,
            size=13 if not body else 11,
            bold=not body,
            color="ink",
        )


def add_freeform_copy(slide, normalized: NormalizedSlide):
    spec = normalized.data
    layout_name = normalized.pattern.layout.name
    if layout_name == "Layout 1":
        title_region, body_region = "rail_title", "rail_body"
    elif layout_name == "Blank White Background":
        title_region, body_region = "blank_title", "blank_body"
    else:
        title_region, body_region = "standard_title", "standard_body"
    title = placeholder_of_types(
        slide, {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
    )
    if title is None and spec.get("title"):
        left, top, width, height = BRAND.region(title_region)
        add_textbox(
            slide,
            left,
            top,
            width,
            height,
            str(spec["title"]),
            font=FONT_TITLE,
            size=28,
            bold=True,
            color="ink",
        )
    if layout_name == "Title Only_Alt" and spec.get("subtitle"):
        left, top, width, height = BRAND.region("title_alt_subhead")
        add_textbox(
            slide,
            left,
            top,
            width,
            height,
            str(spec["subtitle"]),
            font=FONT_TITLE,
            size=16,
            color="ink",
        )
    body = spec.get("body") or spec.get("bullets")
    if body:
        if isinstance(body, list):
            body = "\n".join(
                item if isinstance(item, str) else str(item.get("text") or item.get("title") or "")
                for item in body
            )
        left, top, width, height = BRAND.region(body_region)
        add_textbox(
            slide,
            left,
            top,
            width,
            height,
            str(body),
            font=FONT_BODY,
            size=18,
            color="ink",
        )


def add_timeline(slide, items: list[Any]):
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.75),
        Inches(2.73),
        Inches(8.5),
        Inches(0.03),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = rgb("desk")
    line.line.fill.background()
    step = 8.2 / max(len(items) - 1, 1)
    for i, item in enumerate(items):
        heading, body = item_copy(item)
        x = 0.55 + i * step
        above = i % 2 == 1
        date_y = 2.15 if above else 2.9
        body_y = 1.45 if above else 3.2
        add_textbox(
            slide,
            x,
            date_y,
            1.15,
            0.28,
            heading,
            font=FONT_TITLE,
            size=10,
            bold=True,
            color="water_cooler",
            align=PP_ALIGN.CENTER,
        )
        if body:
            add_textbox(
                slide,
                x,
                body_y,
                1.15,
                0.65,
                body,
                font=FONT_BODY,
                size=9,
                color="ink",
                align=PP_ALIGN.CENTER,
            )


def add_notes(slide, notes: str | None):
    if not notes:
        return
    slide.notes_slide.notes_text_frame.text = notes


def build_slide(prs: Presentation, normalized: NormalizedSlide, footer: str):
    spec = normalized.data
    pattern = normalized.pattern
    layout = find_layout(prs, pattern.layout.name)
    slide = prs.slides.add_slide(layout)

    fill_footer(slide, spec.get("footer", footer), pattern.layout)

    if pattern.composition == Composition.TITLE:
        fill_title(slide, spec.get("title"))
        fill_subtitle(slide, spec.get("subtitle"))
        add_cover_date(slide, spec.get("date"))
    elif pattern.composition == Composition.INTRO:
        fill_intro(slide, normalized)
    elif pattern.composition == Composition.CONTENT:
        fill_title(slide, spec.get("title"))
        body = body_placeholder(slide)
        if body is None:
            raise ValueError("Content layout is missing its body placeholder")
        set_bullets(body, spec.get("bullets") or [])
    elif pattern.composition == Composition.FREEFORM:
        fill_title(slide, spec.get("title"))
        add_freeform_copy(slide, normalized)
    elif pattern.composition == Composition.FEATURE_GRID:
        fill_title(slide, spec.get("title"))
        add_feature_grid(slide, spec)
    elif pattern.composition == Composition.TWO_COLUMN:
        fill_title(slide, spec.get("title"))
        title = placeholder_of_types(
            slide, {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
        )
        if title is not None and title.has_text_frame:
            try:
                title.text_frame._txBody.bodyPr.set("anchor", "ctr")
            except Exception:
                pass
        items = list(spec.get("items") or [])
        add_two_column_items(slide, items, compact=len(items) > 3)
    elif pattern.composition == Composition.SECTION:
        add_section_copy(slide, spec)
    elif pattern.composition == Composition.AGENDA:
        add_agenda(slide, spec)
    elif pattern.composition == Composition.PHOTOS:
        fill_title(slide, spec.get("title"))
        fill_picture_slots(
            slide,
            pattern.layout.picture_slots,
            list(spec.get("images") or []),
        )
    elif pattern.composition == Composition.PHOTO_RIGHT:
        fill_title(slide, spec.get("title"))
        fill_picture_slots(slide, pattern.layout.picture_slots, [spec.get("image")])
        add_photo_right_items(slide, list(spec.get("items") or []))
    elif pattern.composition == Composition.CLOSING:
        add_closing_caption(slide, spec.get("title") or spec.get("caption"))
    elif pattern.composition == Composition.TABLE:
        fill_title(slide, spec.get("title"))
        add_table_slide(slide, spec)
    elif pattern.composition == Composition.CHART:
        fill_title(slide, spec.get("title"))
        add_chart_slide(slide, spec)
    elif pattern.composition == Composition.TIMELINE:
        fill_title(slide, spec.get("title"))
        add_timeline(slide, list(spec.get("items") or []))

    add_notes(slide, spec.get("notes"))
    return slide


def apply_core_props(prs: Presentation, outline: dict[str, Any]):
    title = outline.get("title") or "Workday presentation"
    prs.core_properties.title = title
    prs.core_properties.author = outline.get("author") or "Workday"
    prs.core_properties.last_modified_by = "Workday PPTX skill"
    if outline.get("subject"):
        prs.core_properties.subject = outline["subject"]


def build(outline: dict[str, Any], template: Path, dest: Path):
    prs = Presentation(str(template))
    if len(prs.slides) != 0:
        print("warning: template already contains slides; new slides will be appended.", file=sys.stderr)
    footer = outline.get("footer", "Workday Confidential")
    try:
        slides = BRAND.normalize_outline(outline)
    except ValueError as exc:
        raise SystemExit(str(exc)) from exc
    for normalized in slides:
        build_slide(prs, normalized, footer)
    apply_core_props(prs, outline)
    dest.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(dest))
    print(f"Wrote {dest} ({len(prs.slides)} slides) from {template.name}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build a Workday-branded PPTX from an outline.")
    parser.add_argument("outline", type=Path, help="JSON or YAML outline")
    parser.add_argument("-o", "--output", type=Path, default=None, help="Output .pptx path")
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE, help="Masters-only template")
    args = parser.parse_args(argv)
    if not args.template.is_file():
        raise SystemExit(f"Template not found: {args.template}")
    try:
        outline = load_outline(args.outline)
    except (ValueError, RuntimeError) as exc:
        raise SystemExit(str(exc)) from exc
    dest = args.output
    if dest is None:
        dest = args.outline.with_suffix(".pptx")
    build(outline, args.template, dest)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
