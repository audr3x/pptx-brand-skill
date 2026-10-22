#!/usr/bin/env python3
"""Read back the visible text of a built deck and list outline text that is not on the slides.

Usage:
    python3 scripts/readback.py outline.yaml out/deck.pptx
    python3 scripts/readback.py outline.yaml out/slides/fills.json
    python3 scripts/readback.py outline.yaml out/deck.pptx --show

Run it after every build and before telling the user what the deck contains. It reads
what a viewer sees (slide text, table cells, chart labels; not speaker notes) and checks
every title, bullet, item, cell, and chart label the outline asks for. Exit 1 when
anything is missing, so text dropped by a YAML slip, a slide limit, or a wrong layout
is caught before anyone reports it as done. --show prints the visible text per slide.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from outline_rules import drawn_strings, load, prepare  # noqa: E402


def norm(text: str) -> str:
    text = str(text).replace("\u2019", "'").replace("\u2018", "'").replace("\u201c", '"').replace("\u201d", '"')
    return re.sub(r"\s+", " ", text).strip().lower()


def pptx_visible_text(path: Path) -> list[str]:
    from pptx import Presentation

    texts = []
    for slide in Presentation(str(path)).slides:
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                parts.extend(c.text for row in shape.table.rows for c in row.cells)
            if getattr(shape, "has_chart", False) and shape.has_chart:
                for plot in shape.chart.plots:
                    parts.extend(str(c) for c in plot.categories)
                    for s in plot.series:
                        parts.append(str(s.name))
                        parts.extend(f"{v:g}" for v in s.values if v is not None)
        texts.append(" ".join(parts))
    return texts


def slides_visible_text(path: Path) -> list[str]:
    """Text inserted on each slide of a Google Slides fills batch (fills.json)."""
    data = json.loads(Path(path).read_text())
    texts: dict[int, list[str]] = {}
    for req in data["requests"] if isinstance(data, dict) else data:
        ins = req.get("insertText")
        m = re.match(r"wd_s(\d+)", ins["objectId"]) if ins else None
        if m:
            texts.setdefault(int(m.group(1)), []).append(ins["text"])
    return [" ".join(texts.get(i, [])) for i in range(1, max(texts, default=0) + 1)]


def readback(outline_path: Path, built: Path) -> tuple[list[str], list[str], int]:
    """Return (missing items, visible text per slide, items checked)."""
    target = "pptx" if built.suffix.lower() == ".pptx" else "slides"
    outline, errors, _ = prepare(load(outline_path), target, str(outline_path))
    if errors:
        raise SystemExit(f"{outline_path}: outline has errors; run scripts/check_outline.py first")
    visible = pptx_visible_text(built) if target == "pptx" else slides_visible_text(built)
    slides = outline["slides"]
    missing = []
    # A Slides fills batch has no requests for a slide with no text, so its slide count
    # can be short without anything dropped; the per-item check below still catches drops.
    if target == "pptx" and len(visible) < len(slides):
        missing.append(f"deck has {len(visible)} slides, outline has {len(slides)}")
    checked = 0
    for i, spec in enumerate(slides, 1):
        if not isinstance(spec, dict):
            continue
        seen = norm(visible[i - 1]) if i <= len(visible) else ""
        for field, text in drawn_strings(spec, target):
            checked += 1
            if norm(text) not in seen:
                missing.append(f"slide {i} {field}: {text[:90]!r}")
    return missing, visible, checked


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("outline", type=Path)
    ap.add_argument("built", type=Path, help="deck.pptx, or the Slides fills.json")
    ap.add_argument("--show", action="store_true", help="print the visible text of each slide")
    args = ap.parse_args(argv)
    missing, visible, checked = readback(args.outline, args.built)
    if args.show:
        for i, text in enumerate(visible, 1):
            print(f"slide {i}: {re.sub(r'\s+', ' ', text).strip()}")
    for m in missing:
        print(f"missing: {m}")
    if missing:
        print(f"FAIL: {len(missing)} of {checked} outline items are not on the slides. Fix the outline "
              "(quote text with ': ', split long slides, pick a layout that draws the field) and rebuild")
        return 1
    print(f"ok: all {checked} outline items are on the {len(visible)} slides")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
