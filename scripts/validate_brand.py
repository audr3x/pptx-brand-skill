#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import re
import sys
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

from pptx import Presentation

from brand import BRAND, ROOT


THEME_NAMESPACE = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main"
}
MARKDOWN_LINK = re.compile(r"\[[^\]]+\]\(([^)]+)\)")
HEX_LITERAL = re.compile(r"""["']([0-9A-Fa-f]{6})["']""")
DOCUMENTED_HEX = re.compile(r"#([0-9A-Fa-f]{6})")


def theme_colors(path: Path) -> dict[str, str]:
    with ZipFile(path) as package:
        root = ElementTree.fromstring(package.read("ppt/theme/theme1.xml"))
    scheme = root.find(".//a:clrScheme", THEME_NAMESPACE)
    if scheme is None:
        raise ValueError(f"{path} has no color scheme")
    colors: dict[str, str] = {}
    for entry in scheme:
        value = list(entry)[0].attrib.get("val") or list(entry)[0].attrib.get(
            "lastClr"
        )
        colors[entry.tag.rsplit("}", 1)[-1]] = value
    return colors


def layout_placeholders(
    path: Path,
) -> dict[str, dict[int, tuple[float, float, float, float]]]:
    presentation = Presentation(str(path))
    return {
        layout.name: {
            placeholder.placeholder_format.idx: (
                placeholder.left / 914400,
                placeholder.top / 914400,
                placeholder.width / 914400,
                placeholder.height / 914400,
            )
            for placeholder in layout.placeholders
        }
        for layout in presentation.slide_layouts
    }


def local_links(path: Path) -> list[Path]:
    links: list[Path] = []
    for target in MARKDOWN_LINK.findall(path.read_text()):
        if target.startswith(("http://", "https://", "#")):
            continue
        clean = target.split("#", 1)[0]
        if clean:
            links.append((path.parent / clean).resolve())
    return links


def validate_brand(template: Path, source: Path | None = None) -> list[str]:
    errors: list[str] = []
    presentation = Presentation(str(template))
    if presentation.slides:
        errors.append("assets/template.pptx must contain no slides")
    actual_layouts = layout_placeholders(template)
    expected_layouts = set(BRAND.layouts)
    if set(actual_layouts) != expected_layouts:
        errors.append(
            "Template layouts differ from the registry: "
            f"template={sorted(actual_layouts)}, registry={sorted(expected_layouts)}"
        )
    for layout in BRAND.layouts.values():
        declared = set(layout.placeholders.values())
        for slot in layout.speaker_slots:
            declared.update(slot.values())
        declared.update(layout.picture_slots)
        actual = actual_layouts.get(layout.name, {})
        missing = declared - set(actual)
        if missing:
            errors.append(
                f"{layout.name} references missing placeholders {sorted(missing)}"
            )
        picture_indexes = (
            tuple(slot["picture"] for slot in layout.speaker_slots)
            or layout.picture_slots
        )
        for idx, expected in zip(picture_indexes, layout.picture_regions):
            if idx not in actual:
                continue
            if any(abs(a - b) > 0.02 for a, b in zip(actual[idx], expected)):
                errors.append(
                    f"{layout.name} picture placeholder {idx} moved from {expected}"
                )
    actual_theme = theme_colors(template)
    for theme_slot, token in BRAND.raw["theme_map"].items():
        expected = BRAND.colors[token]
        actual = actual_theme.get(theme_slot)
        if actual != expected:
            errors.append(
                f"Theme {theme_slot} is {actual}; registry expects {expected}"
            )
    represented_layouts = {pattern.layout.name for pattern in BRAND.patterns.values()}
    missing_patterns = expected_layouts - represented_layouts
    if missing_patterns:
        errors.append(
            f"Layouts without an outline pattern: {sorted(missing_patterns)}"
        )
    for script_name in ("build_deck.py", "build_slides.py"):
        script = ROOT / "scripts" / script_name
        literals = sorted(set(HEX_LITERAL.findall(script.read_text())))
        if literals:
            errors.append(f"{script_name} duplicates hex colors: {literals}")
    documented = {
        value.upper()
        for value in DOCUMENTED_HEX.findall((ROOT / "BRAND.md").read_text())
    }
    allowed_documented = {
        value.upper() for value in BRAND.colors.values()
    } | {"9FCFFF", "A4A3A4", "F38B00"}
    unexpected = documented - allowed_documented
    if unexpected:
        errors.append(f"BRAND.md contains unknown colors: {sorted(unexpected)}")
    for skill in (
        ROOT / ".cursor" / "skills" / "workday-pptx" / "SKILL.md",
        ROOT / ".cursor" / "skills" / "workday-google-slides" / "SKILL.md",
    ):
        text = skill.read_text()
        if not text.startswith("---\n") or "\nname:" not in text or "\ndescription:" not in text:
            errors.append(f"{skill} needs name and description frontmatter")
        for link in local_links(skill):
            if not link.exists():
                errors.append(f"{skill} links to missing file {link}")
    if source is not None:
        if theme_colors(source) != actual_theme:
            errors.append("The source deck theme differs from assets/template.pptx")
        with ZipFile(source) as package:
            source_xml = "\n".join(
                package.read(name).decode("utf-8", "ignore")
                for name in package.namelist()
                if name.endswith(".xml")
            ).upper()
        for name, value in BRAND.colors.items():
            if value.upper() not in source_xml:
                errors.append(f"Source deck does not contain {name} #{value}")
        expected_hash = BRAND.raw["provenance"].get("source_sha256")
        actual_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        if expected_hash and actual_hash != expected_hash:
            errors.append(
                "Source deck checksum differs from the recorded extraction source"
            )
    return errors


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Validate the Workday brand system")
    parser.add_argument(
        "--template",
        type=Path,
        default=ROOT / "assets" / "template.pptx",
    )
    parser.add_argument("--source", type=Path)
    args = parser.parse_args(argv)
    errors = validate_brand(args.template, args.source)
    if errors:
        for error in errors:
            print(f"error: {error}", file=sys.stderr)
        return 1
    print(
        f"OK: {len(BRAND.colors)} colors, {len(BRAND.layouts)} layouts, "
        f"{len(BRAND.patterns)} outline patterns"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
